#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""Assert fosslight_dependency Excel reports under a result tree are non-empty.

Looks for ``fosslight_report_dep_*.xlsx`` and requires ``DEP_FL_Dependency``
to have at least 2 non-empty rows (header + data).

Use ``--profile`` to select which result dirs must exist for each OS tox env.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import load_workbook

DEP_SHEET_NAME = "DEP_FL_Dependency"
MIN_ROWS = 2

REQUIRED_BY_PROFILE = {
    "ubuntu": frozenset({
        "android",
        "cargo",
        "exclude",
        "gradle",
        "helm",
        "maven1",
        "maven2",
        "mod",
        "multi_pypi_npm",
        "npm1",
        "npm2",
        "nuget1",
        "nuget2",
        "pub",
        "pypi",
    }),
    # Windows pytest markers cover a subset of managers (no helm/npm/cocoapods).
    "windows": frozenset({
        "android",
        "cargo",
        "exclude",
        "gradle",
        "maven2",
        "mod",
        "nuget1",
        "nuget2",
        "pub",
        "pypi",
    }),
    # macOS pytest -m macos currently covers CocoaPods + pub.
    "macos": frozenset({
        "cocoapods",
        "pub",
    }),
}


def _row_count(sheet) -> int:
    count = 0
    for row in sheet.iter_rows(values_only=True):
        if any(cell is not None and str(cell).strip() != "" for cell in row):
            count += 1
    return count


def check_report(path: Path) -> tuple[bool, str]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if DEP_SHEET_NAME not in workbook.sheetnames:
            return False, f"missing sheet {DEP_SHEET_NAME}; sheets={workbook.sheetnames}"
        rows = _row_count(workbook[DEP_SHEET_NAME])
        if rows < MIN_ROWS:
            return False, f"{DEP_SHEET_NAME} has {rows} row(s); expected >= {MIN_ROWS}"
        return True, f"{DEP_SHEET_NAME} rows={rows}"
    finally:
        workbook.close()


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "result_root",
        type=Path,
        help="Root directory that contains per-manager result folders (e.g. tests/result)",
    )
    parser.add_argument(
        "--profile",
        choices=sorted(REQUIRED_BY_PROFILE),
        default="ubuntu",
        help="Which tox OS profile's required result dirs to enforce (default: ubuntu)",
    )
    parser.add_argument(
        "--all",
        action="store_true",
        help="Check every fosslight_report_dep_*.xlsx (ignore --profile allowlist)",
    )
    args = parser.parse_args()
    root: Path = args.result_root
    required_dirs = REQUIRED_BY_PROFILE[args.profile]

    if not root.is_dir():
        print(f"ERROR: result root not found: {root}", file=sys.stderr)
        return 2

    reports = sorted(root.glob("**/fosslight_report_dep_*.xlsx"))
    if not reports:
        print(f"ERROR: no fosslight_report_dep_*.xlsx under {root}", file=sys.stderr)
        return 1

    if not args.all:
        filtered = []
        found_tops: set[str] = set()
        for report in reports:
            try:
                rel = report.relative_to(root)
            except ValueError:
                filtered.append(report)
                continue
            top = rel.parts[0] if rel.parts else ""
            if top in required_dirs:
                filtered.append(report)
                found_tops.add(top)
            else:
                print(f"[SKIP] {rel}: not in required {args.profile} result dirs")
        reports = filtered
        missing = sorted(required_dirs - found_tops)
        if missing:
            print(
                f"ERROR: missing required {args.profile} result dir(s): {', '.join(missing)}",
                file=sys.stderr,
            )
            return 1

    if not reports:
        print(f"ERROR: no required fosslight_report_dep_*.xlsx under {root}", file=sys.stderr)
        return 1

    failed = 0
    for report in reports:
        ok, message = check_report(report)
        status = "OK" if ok else "FAIL"
        try:
            rel = report.relative_to(root)
        except ValueError:
            rel = report
        print(f"[{status}] {rel}: {message}")
        if not ok:
            failed += 1

    if failed:
        print(f"ERROR: {failed}/{len(reports)} report(s) failed DEP sheet check", file=sys.stderr)
        return 1

    print(
        f"SUCCESS: {len(reports)} report(s) have non-empty {DEP_SHEET_NAME} "
        f"(profile={args.profile})"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
