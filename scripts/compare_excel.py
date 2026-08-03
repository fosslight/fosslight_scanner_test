#!/usr/bin/env python3
# SPDX-License-Identifier: Apache-2.0
"""Compare two FOSSLight Report Excel files and print differences."""

from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

# Columns that are expected to differ between independent runs / install methods.
IGNORE_COLUMNS = {
    "id",
    "tlsh",  # fuzzy hash can vary slightly across environments
}

# Comma-separated multi-value columns: compare as unordered sets.
LIST_COLUMNS = {
    "license",
    "depends on",
}

# Scanner Info keys that always differ between independent runs.
IGNORE_SCANNER_INFO_KEYS = {
    "running time",
    "analyzed path",
}

COVER_SHEET_NAME = "Scanner Info"

# Key columns used to align rows within a sheet (first match wins).
KEY_COLUMNS = (
    "source path",
    "binary path",
    "package url",
    "oss name",
)


def _norm(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _norm_cell(column: str, value: Any) -> str:
    """Normalize a cell; sort comma-separated lists so order-only diffs are ignored."""
    text = _norm(value)
    if not text:
        return ""
    if column.lower() in LIST_COLUMNS:
        parts = [p.strip() for p in text.split(",") if p.strip()]
        return ",".join(sorted(parts, key=str.casefold))
    return text


def _header_map(header_row: tuple[Any, ...]) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        name = _norm(cell)
        if name:
            mapping[name.lower()] = idx
    return mapping


def _row_key(headers: dict[str, int], row: tuple[Any, ...]) -> str:
    for key_name in KEY_COLUMNS:
        if key_name in headers:
            value = _norm(row[headers[key_name]]) if headers[key_name] < len(row) else ""
            if value:
                return f"{key_name}={value}"
    # Fallback: join all non-ignored cells
    parts = []
    for name, idx in sorted(headers.items(), key=lambda x: x[1]):
        if name in IGNORE_COLUMNS:
            continue
        if idx < len(row):
            parts.append(_norm(row[idx]))
    return "|".join(parts)


def _sheet_rows(ws) -> tuple[dict[str, int], list[tuple[Any, ...]]]:
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return {}, []
    headers = _header_map(rows[0])
    data = [r for r in rows[1:] if any(_norm(c) for c in r)]
    return headers, data


def _scanner_info_map(ws) -> dict[str, str]:
    """Parse Scanner Info sheet as key(A)/value(B) pairs."""
    mapping: dict[str, str] = {}
    for row in ws.iter_rows(values_only=True):
        if not row:
            continue
        key = _norm(row[0] if len(row) > 0 else "")
        if not key or key.lower() == "about the scanner":
            continue
        value = _norm(row[1] if len(row) > 1 else "")
        mapping[key] = value
    return mapping


def _compare_scanner_info(sheet_name: str, left_ws, right_ws) -> list[dict[str, Any]]:
    diffs: list[dict[str, Any]] = []
    left = _scanner_info_map(left_ws)
    right = _scanner_info_map(right_ws)
    keys = sorted(set(left) | set(right), key=str.lower)
    for key in keys:
        if key.lower() in IGNORE_SCANNER_INFO_KEYS:
            continue
        lv = left.get(key, "")
        rv = right.get(key, "")
        if lv != rv:
            diffs.append(
                {
                    "sheet": sheet_name,
                    "type": "scanner_info_changed",
                    "key": key,
                    "pypi": lv,
                    "github": rv,
                }
            )
    return diffs


def _compare_sheet(sheet_name: str, left_ws, right_ws) -> list[dict[str, Any]]:
    diffs: list[dict[str, Any]] = []
    if sheet_name == COVER_SHEET_NAME:
        return _compare_scanner_info(sheet_name, left_ws, right_ws)

    left_headers, left_rows = _sheet_rows(left_ws)
    right_headers, right_rows = _sheet_rows(right_ws)

    # Sheets without recognizable headers — compare line-wise.
    if not left_headers and not right_headers:
        left_all = list(left_ws.iter_rows(values_only=True))
        right_all = list(right_ws.iter_rows(values_only=True))
        max_len = max(len(left_all), len(right_all))
        for i in range(max_len):
            lrow = left_all[i] if i < len(left_all) else ()
            rrow = right_all[i] if i < len(right_all) else ()
            if tuple(_norm(c) for c in lrow) != tuple(_norm(c) for c in rrow):
                diffs.append(
                    {
                        "sheet": sheet_name,
                        "type": "row_changed",
                        "row": i + 1,
                        "pypi": list(lrow),
                        "github": list(rrow),
                    }
                )
        return diffs

    common_cols = sorted(set(left_headers) & set(right_headers) - IGNORE_COLUMNS)
    only_left_cols = sorted(set(left_headers) - set(right_headers) - IGNORE_COLUMNS)
    only_right_cols = sorted(set(right_headers) - set(left_headers) - IGNORE_COLUMNS)
    for col in only_left_cols:
        diffs.append({"sheet": sheet_name, "type": "column_removed", "column": col})
    for col in only_right_cols:
        diffs.append({"sheet": sheet_name, "type": "column_added", "column": col})

    left_map: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    right_map: dict[str, list[tuple[Any, ...]]] = defaultdict(list)
    for row in left_rows:
        left_map[_row_key(left_headers, row)].append(row)
    for row in right_rows:
        right_map[_row_key(right_headers, row)].append(row)

    all_keys = sorted(set(left_map) | set(right_map))
    for key in all_keys:
        lrows = left_map.get(key, [])
        rrows = right_map.get(key, [])
        if not rrows:
            diffs.append(
                {
                    "sheet": sheet_name,
                    "type": "row_only_in_pypi",
                    "key": key,
                    "pypi": [_row_as_dict(left_headers, r) for r in lrows],
                }
            )
            continue
        if not lrows:
            diffs.append(
                {
                    "sheet": sheet_name,
                    "type": "row_only_in_github",
                    "key": key,
                    "github": [_row_as_dict(right_headers, r) for r in rrows],
                }
            )
            continue

        # Compare first matching occurrence; extras treated as add/remove.
        pair_count = min(len(lrows), len(rrows))
        for i in range(pair_count):
            cell_diffs = []
            for col in common_cols:
                lv = (
                    _norm_cell(col, lrows[i][left_headers[col]])
                    if left_headers[col] < len(lrows[i])
                    else ""
                )
                rv = (
                    _norm_cell(col, rrows[i][right_headers[col]])
                    if right_headers[col] < len(rrows[i])
                    else ""
                )
                if lv != rv:
                    cell_diffs.append({"column": col, "pypi": lv, "github": rv})
            if cell_diffs:
                diffs.append(
                    {
                        "sheet": sheet_name,
                        "type": "cell_changed",
                        "key": key,
                        "changes": cell_diffs,
                    }
                )
        for extra in lrows[pair_count:]:
            diffs.append(
                {
                    "sheet": sheet_name,
                    "type": "row_only_in_pypi",
                    "key": key,
                    "pypi": [_row_as_dict(left_headers, extra)],
                }
            )
        for extra in rrows[pair_count:]:
            diffs.append(
                {
                    "sheet": sheet_name,
                    "type": "row_only_in_github",
                    "key": key,
                    "github": [_row_as_dict(right_headers, extra)],
                }
            )
    return diffs


def _row_as_dict(headers: dict[str, int], row: tuple[Any, ...]) -> dict[str, str]:
    result = {}
    for name, idx in headers.items():
        if name in IGNORE_COLUMNS:
            continue
        result[name] = _norm(row[idx]) if idx < len(row) else ""
    return result


def compare_excels(pypi_path: Path, github_path: Path) -> list[dict[str, Any]]:
    left = load_workbook(pypi_path, data_only=True)
    right = load_workbook(github_path, data_only=True)

    diffs: list[dict[str, Any]] = []
    left_sheets = set(left.sheetnames)
    right_sheets = set(right.sheetnames)

    for name in sorted(left_sheets - right_sheets):
        diffs.append({"sheet": name, "type": "sheet_only_in_pypi"})
    for name in sorted(right_sheets - left_sheets):
        diffs.append({"sheet": name, "type": "sheet_only_in_github"})

    for name in sorted(left_sheets & right_sheets):
        diffs.extend(_compare_sheet(name, left[name], right[name]))
    return diffs


def _escape_md(value: Any, max_len: int = 120) -> str:
    text = "" if value is None else str(value)
    text = text.replace("\n", " ").replace("|", "\\|").replace("`", "'")
    if len(text) > max_len:
        return text[: max_len - 1] + "…"
    return text


def _flatten_diff_rows(diffs: list[dict[str, Any]]) -> list[dict[str, str]]:
    """Flatten nested diffs into table rows: sheet/type/key/column/pypi/github."""
    rows: list[dict[str, str]] = []
    for diff in diffs:
        dtype = diff.get("type", "")
        sheet = str(diff.get("sheet", ""))
        key = str(diff.get("key", diff.get("column", diff.get("row", ""))))

        if dtype == "cell_changed":
            for change in diff.get("changes", []):
                rows.append(
                    {
                        "sheet": sheet,
                        "type": dtype,
                        "key": key,
                        "column": str(change.get("column", "")),
                        "pypi": str(change.get("pypi", "")),
                        "github": str(change.get("github", "")),
                    }
                )
        elif dtype == "scanner_info_changed":
            rows.append(
                {
                    "sheet": sheet,
                    "type": dtype,
                    "key": key,
                    "column": key,
                    "pypi": str(diff.get("pypi", "")),
                    "github": str(diff.get("github", "")),
                }
            )
        elif dtype in ("row_only_in_pypi", "row_only_in_github"):
            side = "pypi" if dtype == "row_only_in_pypi" else "github"
            detail = json.dumps(diff.get(side), ensure_ascii=False)
            rows.append(
                {
                    "sheet": sheet,
                    "type": dtype,
                    "key": key,
                    "column": "(row)",
                    "pypi": detail if side == "pypi" else "",
                    "github": detail if side == "github" else "",
                }
            )
        elif dtype == "row_changed":
            rows.append(
                {
                    "sheet": sheet,
                    "type": dtype,
                    "key": f"row={diff.get('row')}",
                    "column": "(row)",
                    "pypi": str(diff.get("pypi", "")),
                    "github": str(diff.get("github", "")),
                }
            )
        elif dtype in ("column_added", "column_removed", "sheet_only_in_pypi", "sheet_only_in_github"):
            rows.append(
                {
                    "sheet": sheet,
                    "type": dtype,
                    "key": key,
                    "column": str(diff.get("column", "")),
                    "pypi": "present" if "pypi" in dtype or dtype == "column_removed" else "",
                    "github": "present" if "github" in dtype or dtype == "column_added" else "",
                }
            )
        else:
            rows.append(
                {
                    "sheet": sheet,
                    "type": dtype,
                    "key": key,
                    "column": "",
                    "pypi": json.dumps(diff, ensure_ascii=False),
                    "github": "",
                }
            )
    return rows


def format_diffs_markdown(diffs: list[dict[str, Any]]) -> str:
    if not diffs:
        return "No differences found between PyPI and GitHub Excel reports.\n"

    rows = _flatten_diff_rows(diffs)
    lines = [
        f"Found **{len(rows)}** difference cell(s) "
        f"(from {len(diffs)} diff group(s)).",
        "",
        "| # | Sheet | Type | Key | Column | PyPI | GitHub |",
        "|---|-------|------|-----|--------|------|--------|",
    ]
    for i, row in enumerate(rows, 1):
        lines.append(
            "| {num} | {sheet} | {dtype} | {key} | {column} | {pypi} | {github} |".format(
                num=i,
                sheet=_escape_md(row["sheet"], 40),
                dtype=_escape_md(row["type"], 40),
                key=_escape_md(row["key"], 60),
                column=_escape_md(row["column"], 30),
                pypi=_escape_md(row["pypi"]),
                github=_escape_md(row["github"]),
            )
        )
    lines.append("")
    return "\n".join(lines)


def _print_diffs(diffs: list[dict[str, Any]]) -> None:
    print(format_diffs_markdown(diffs), end="")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("pypi_excel", type=Path, help="Excel from pip install fosslight_scanner")
    parser.add_argument("github_excel", type=Path, help="Excel from GitHub package install")
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Optional JSON path to write the full diff list",
    )
    parser.add_argument(
        "--md",
        type=Path,
        help="Optional Markdown path to write the diff table",
    )
    args = parser.parse_args()

    if not args.pypi_excel.is_file():
        print(f"ERROR: PyPI excel not found: {args.pypi_excel}", file=sys.stderr)
        return 2
    if not args.github_excel.is_file():
        print(f"ERROR: GitHub excel not found: {args.github_excel}", file=sys.stderr)
        return 2

    diffs = compare_excels(args.pypi_excel, args.github_excel)
    md_text = format_diffs_markdown(diffs)
    print(md_text, end="")

    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(json.dumps(diffs, indent=2, ensure_ascii=False), encoding="utf-8")
        print(f"Wrote JSON diff: {args.output}")

    md_path = args.md
    if md_path is None and args.output:
        md_path = args.output.with_suffix(".md")
    if md_path is not None:
        md_path.parent.mkdir(parents=True, exist_ok=True)
        md_path.write_text(md_text, encoding="utf-8")
        print(f"Wrote Markdown diff: {md_path}")

    return 1 if diffs else 0


if __name__ == "__main__":
    sys.exit(main())
